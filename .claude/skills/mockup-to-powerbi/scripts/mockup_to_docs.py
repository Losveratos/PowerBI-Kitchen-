#!/usr/bin/env python3
"""mockup_to_docs.py — Workshop-Doku und PowerPoint aus der MockupKitchen-Spec.

Zweiter Ausgang des Skills `mockup-to-powerbi`: nicht der Bericht, sondern das
Protokoll. Aus `mockup-spec.json` (specVersion 1 oder 2) entstehen

  WORKSHOP-DOKU.md    gleiche Struktur wie `buildDocs` in assets/mockup/export.js
                      (wird übersprungen, wenn die Datei schon neben der Spec
                      liegt — das Tool schreibt sie selbst; mit --force-md
                      trotzdem neu erzeugen)
  WORKSHOP-DOKU.pptx  Folien: Titel · Gestaltungsentscheidungen · je Seite ein
                      maßstäbliches Wireframe aus nativen Shapes + eine
                      Kachel-Tabelle · Navigation/Drill · neue Kennzahlen ·
                      offene Punkte · nächste Schritte

    python mockup_to_docs.py mockup-spec.json [--out .] [--no-pptx] [--force-md]
                             [--docs WORKSHOP-DOKU.md]

Die PowerPoint braucht `python-pptx`:

    pip install python-pptx

Fehlt das Paket, schreibt das Skript trotzdem die Markdown-Fassung und meldet
den fehlenden Import. Alternative ohne Installation: die Markdown-Datei an den
Skill `anthropic-skills:pptx` geben.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

# --------------------------------------------------------------------------- #
# Rollen-Labels aus assets/mockup/catalog.js (R). Nur zur Anzeige.
# --------------------------------------------------------------------------- #
ROLE_LABEL = {
    "category": "Kategorie / Zeit",
    "subcategory": "Unterkategorie",
    "series": "Reihe / Legende",
    "ac": "AC · Ist-Wert",
    "ref": "Referenz (PL / PY / BU)",
    "fc": "FC-Flag (1/0)",
    "values": "Werte",
    "rows": "Zeilen",
    "columns": "Spalten",
    "x": "X-Wert", "y": "Y-Wert", "size": "Größe",
    "indicator": "Kennzahl", "goal": "Ziel / Referenz",
    "start": "Start", "end": "Ende",
    "field": "Feld", "text": "Text",
}
ENGINE_LABEL = {"ck": "ChartKitchen", "native": "Nativ", "deneb": "Deneb"}
TILE_STYLE = {"border": "mit feinem Rahmen", "shadow": "mit weichem Schatten",
              "flat": "flach"}
FILTER_MODE = {"right": "als Panel rechts", "left": "als Panel links",
               "top": "als Leiste oben", "burger": "als Burger-Menü (Bookmark)"}
PRESET_LABEL = {"1280x720": "HD", "1920x1080": "Full HD", "3840x2160": "Ultra HD"}

INK = "0F1E2E"
INK_SOFT = "475569"
LINE = "D5D8DC"
PAPER = "FFFFFF"


def slug(text: str, limit: int = 40) -> str:
    s = unicodedata.normalize("NFD", str(text or "seite"))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")[:limit]
    return s or "Seite"


def role_label(key: str) -> str:
    return ROLE_LABEL.get(key, key)


# --------------------------------------------------------------------------- #
# Spec normalisieren (identisch zur Logik in mockup_to_pbir.py)
# --------------------------------------------------------------------------- #
DEFAULT_DESIGN = {
    "cornerRadius": 0, "tileStyle": "border", "pageBackground": "#F4F4F1",
    "tileBackground": "#FFFFFF", "headerStyle": "dark", "accent": "#C25A2D",
    "fontScale": 1.0,
}


def normalize(spec: dict) -> dict:
    meta = spec.get("meta") or {}
    version = int(meta.get("specVersion") or (2 if isinstance(spec.get("pages"), list) else 1))
    canvas = spec.get("canvas") or {}
    zones = spec.get("zones") or {}
    ui_scale = float(canvas.get("uiScale") or 1) or 1.0

    design = dict(DEFAULT_DESIGN)
    design.update({k: v for k, v in (spec.get("design") or {}).items() if v is not None})
    design["fontScale"] = float(design.get("fontScale") or ui_scale or 1)

    if version >= 2:
        pages = [{
            "id": p.get("id") or f"p{i + 1}", "index": int(p.get("index") or i + 1),
            "name": p.get("name") or f"Seite {i + 1}", "notes": p.get("notes") or "",
            "contentRect": p.get("contentRect") or zones.get("content") or {},
            "visuals": p.get("visuals") or [],
        } for i, p in enumerate(spec.get("pages") or [])]
        links = spec.get("links") or []
    else:
        name = (spec.get("page") or {}).get("name") or "Seite"
        pages = [{"id": "p1", "index": 1, "name": name,
                  "notes": (spec.get("page") or {}).get("notes") or "",
                  "contentRect": zones.get("content") or {},
                  "visuals": spec.get("visuals") or []}]
        links = []
        filt = zones.get("filter")
        if filt and not filt.get("mode"):
            filt["mode"] = filt.get("side") or "right"
    return {"version": version, "canvas": canvas, "uiScale": ui_scale,
            "design": design, "zones": zones, "pages": pages, "links": links,
            "meta": meta}


def fields_summary(v: dict) -> str:
    roles = v.get("roles") or {}
    parts = []
    for key, fl in roles.items():
        names = ", ".join(f["name"] + (" (neu)" if f.get("isNew") else "") for f in fl)
        parts.append(f"{role_label(key)}: {names}")
    return "; ".join(parts) or "–"


def open_points(spec: dict, n: dict) -> list[str]:
    out: list[str] = []
    for f in spec.get("newFields") or []:
        if f.get("openQuestion"):
            out.append(f"{f['name']}: {f['openQuestion']}")
    for p in n["pages"]:
        for v in p["visuals"]:
            if "?" in (v.get("notes") or ""):
                out.append(f"{p['name']} / {v.get('title')}: {v['notes']}")
            for w in (v.get("warnings") or []):
                if w.startswith("Pflichtrolle"):
                    out.append(f"{p['name']} / {v.get('title')}: {w}")
    out += list(spec.get("warnings") or [])
    return out


# --------------------------------------------------------------------------- #
# WORKSHOP-DOKU.md  (Struktur wie buildDocs in assets/mockup/export.js)
# --------------------------------------------------------------------------- #
def build_docs_md(spec: dict, n: dict) -> str:
    z, d = n["zones"], n["design"]
    meta = n["meta"]
    canvas = n["canvas"]
    model = spec.get("model") or {}
    L: list[str] = []
    L += [f"# Workshop-Dokumentation · {meta.get('name', 'Mockup')}", "",
          f"Stand: {date.today().isoformat()} · erstellt mit "
          f"{meta.get('tool', 'MockupKitchen byDatenWG')} {meta.get('version', '')}".rstrip(),
          "", "| | |", "|---|---|",
          "| Teilnehmende | [ausfüllen] |",
          "| Ziel des Berichts | [ausfüllen] |",
          "| Zielgruppe | [ausfüllen] |",
          f"| Seiten | {' · '.join(p['name'] for p in n['pages'])} |",
          f"| Datenmodell | {model.get('source') or 'noch nicht angebunden'} "
          f"({len(model.get('tables') or [])} Tabellen) |", ""]

    preset = canvas.get("preset")
    fmt = f"- Format {canvas.get('width')} × {canvas.get('height')} px"
    if preset and preset != "custom":
        fmt += f" ({PRESET_LABEL.get(preset, preset)})"
    header = z.get("header")
    if header:
        h = (f"„{header.get('title', '')}\""
             + (f" · {header['subtitle']}" if header.get("subtitle") else "")
             + f", Stil {header.get('style', d['headerStyle'])}")
        logo = header.get("logoPos") or ("left" if header.get("logo") else "none")
        h += ", Logo " + {"none": "ohne", "right": "rechts"}.get(logo, "links")
    else:
        h = "ohne"
    filt = z.get("filter")
    if filt:
        mode = filt.get("mode") or filt.get("side") or "right"
        f_txt = FILTER_MODE.get(mode, mode)
        sl = filt.get("slicers") or []
        f_txt += (": " + ", ".join(s["name"] for s in sl)) if sl else ", noch ohne Felder"
    else:
        f_txt = "keine"
    L += ["## Entscheidungen zur Gestaltung", "", fmt + ".",
          f"- Kopfband {h}.",
          f"- Filter {f_txt}.",
          f"- Fußleiste " + (f"„{z['footer'].get('text', '')}\"." if z.get("footer") else "ohne."),
          f"- Kacheln " + (f"gerundet ({d['cornerRadius']} px)" if d["cornerRadius"] else "eckig")
          + f", {TILE_STYLE.get(d['tileStyle'], d['tileStyle'])}, Seitenhintergrund "
          f"{d['pageBackground']}, Akzent {d['accent']}.", ""]

    for p in n["pages"]:
        L += [f"## Seite {p['index']} · {p['name']}", "",
              f"**Zweck:** {p['notes']}" if p["notes"] else "**Zweck:** [ausfüllen]", ""]
        if p["visuals"]:
            L += ["| # | Kachel | Darstellung | Felder | Notizen |", "|---|---|---|---|---|"]
            for i, v in enumerate(p["visuals"], start=1):
                render = (v.get("label") or v.get("kind", ""))
                if v.get("scenario"):
                    render += f", {v['scenario']}"
                if v.get("engine") == "ck":
                    render += ", ChartKitchen"
                note = " · ".join(x for x in [v.get("notes") or "",
                                              f"→ {v['link']['pageName']}" if v.get("link") else ""] if x)
                title = v.get("title", "")
                if v.get("subtitle"):
                    title += f" ({v['subtitle']})"
                L.append(f"| {p['index']}.{i} | {title} | {render} | "
                         f"{fields_summary(v)} | {note or '–'} |")
        else:
            L.append("_Noch keine Kacheln._")
        L.append("")

    if n["links"]:
        L += ["## Navigation und Drill-Wege", ""]
        for l in n["links"]:
            kind = "Seitenwechsel" if l.get("kind") == "navigation" else "Drill-through"
            L.append(f"- Von „{l['fromPage']}\" ({l['fromVisual']}) nach "
                     f"„{l['toPage']}\" ({kind})")
        L.append("")

    L += ["## Kennzahlen und Felder", ""]
    used = [f for f in (model.get("usedFields") or []) if not f.get("isNew")]
    L += [("Verwendet aus dem Modell: " + ", ".join(f"`{f['ref']}`" for f in used))
          if used else "Noch keine Modellfelder gebunden.", ""]
    if spec.get("newFields"):
        L += ["### Neu zu erstellen", "",
              "| Feld | Art | Tabelle | Beschreibung / Logik | Offene Frage |",
              "|---|---|---|---|---|"]
        for f in spec["newFields"]:
            L.append(f"| {f['name']} | "
                     f"{'Kennzahl' if f.get('kind') == 'measure' else 'Dimension'} | "
                     f"{f['table']} | {f.get('description') or '[ausfüllen]'} | "
                     f"{f.get('openQuestion') or '–'} |")
        L.append("")

    op = open_points(spec, n)
    L += ["## Offene Punkte", ""]
    L += [f"- [ ] {o}" for o in op] or ["- [ ] keine offenen Punkte erfasst"]
    L += ["", "## Nächste Schritte", "",
          "- [ ] Fehlende Kennzahlen im Semantikmodell anlegen (siehe oben)",
          "- [ ] Seiten mit dem Skill `mockup-to-powerbi` ins PBIP übertragen",
          "- [ ] Review der gebauten Seiten mit den Teilnehmenden", ""]
    return "\n".join(L)


# --------------------------------------------------------------------------- #
# WORKSHOP-DOKU.pptx
# --------------------------------------------------------------------------- #
def hexcolor(value: str, fallback: str = INK):
    v = str(value or "").lstrip("#")
    return v.upper() if re.fullmatch(r"[0-9A-Fa-f]{6}", v or "") else fallback


def build_pptx(spec: dict, n: dict, target: Path) -> tuple[bool, str]:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.util import Emu, Pt
    except ImportError:
        return False, ("python-pptx ist nicht installiert — keine PowerPoint erzeugt.\n"
                       "  Installieren:  pip install python-pptx\n"
                       "  Alternative:   WORKSHOP-DOKU.md an den Skill "
                       "`anthropic-skills:pptx` geben.")

    d = n["design"]
    accent = RGBColor.from_string(hexcolor(d["accent"], "C25A2D"))
    ink = RGBColor.from_string(INK)
    ink_soft = RGBColor.from_string(INK_SOFT)
    line = RGBColor.from_string(LINE)
    page_bg = RGBColor.from_string(hexcolor(d["pageBackground"], "F4F4F1"))
    tile_bg = RGBColor.from_string(hexcolor(d["tileBackground"], "FFFFFF"))

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 16:9
    SW, SH = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]

    def textbox(slide, x, y, w, h, text, *, size=18, bold=False, color=ink,
                align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
        box = slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
        tf = box.text_frame
        tf.word_wrap = wrap
        tf.vertical_anchor = anchor
        lines = str(text).split("\n")
        for i, ln in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            run = para.add_run()
            run.text = ln
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
        return box

    def slide_frame(title: str, kicker: str = "") -> object:
        s = prs.slides.add_slide(blank)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Emu(45720 * 22))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string(PAPER)
        bar.line.color.rgb = line
        bar.line.width = Pt(0.75)
        bar.shadow.inherit = False
        rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Emu(45720 * 22), SW, Emu(22860))
        rule.fill.solid()
        rule.fill.fore_color.rgb = accent
        rule.line.fill.background()
        rule.shadow.inherit = False
        textbox(s, Emu(457200), Emu(150000), SW - Emu(914400), Emu(500000),
                title, size=24, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        if kicker:
            textbox(s, Emu(457200), Emu(150000), SW - Emu(914400), Emu(500000),
                    kicker, size=12, color=ink_soft, align=PP_ALIGN.RIGHT,
                    anchor=MSO_ANCHOR.MIDDLE)
        return s

    def bullets(slide, items, *, x=457200, y=1300000, w=None, size=15, gap=430000):
        w = w or (SW - Emu(914400))
        for i, it in enumerate(items):
            textbox(slide, Emu(x), Emu(y + i * gap), w, Emu(gap), "•  " + it, size=size)

    # ---------------------------------------------------------------- Titel
    s = prs.slides.add_slide(blank)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor.from_string(PAPER)
    band.line.fill.background()
    band.shadow.inherit = False
    stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Emu(2200000), Emu(180000),
                                Emu(1500000))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    stripe.shadow.inherit = False
    textbox(s, Emu(700000), Emu(2150000), SW - Emu(1400000), Emu(700000),
            n["meta"].get("name", "Mockup"), size=40, bold=True)
    textbox(s, Emu(700000), Emu(2900000), SW - Emu(1400000), Emu(500000),
            "Workshop-Dokumentation zur Berichtsskizze", size=20, color=ink_soft)
    model = spec.get("model") or {}
    textbox(s, Emu(700000), Emu(3500000), SW - Emu(1400000), Emu(900000),
            f"{len(n['pages'])} Seite(n) · {sum(len(p['visuals']) for p in n['pages'])} Kacheln"
            f" · {n['canvas'].get('width')} × {n['canvas'].get('height')} px\n"
            f"Datenmodell: {model.get('source') or 'noch nicht angebunden'}\n"
            f"Stand: {date.today().strftime('%d.%m.%Y')} · "
            f"{n['meta'].get('tool', 'MockupKitchen byDatenWG')} "
            f"{n['meta'].get('version', '')}".rstrip(),
            size=14, color=ink_soft)

    # ------------------------------------------------- Gestaltungsentscheidungen
    z = n["zones"]
    header = z.get("header")
    filt = z.get("filter")
    items = [f"Format {n['canvas'].get('width')} × {n['canvas'].get('height')} px"
             + (f" ({PRESET_LABEL.get(n['canvas'].get('preset'), '')})"
                if PRESET_LABEL.get(n["canvas"].get("preset")) else "")
             + f", Skalierung ×{n['uiScale']}"]
    if header:
        logo = header.get("logoPos") or ("left" if header.get("logo") else "none")
        items.append(f"Kopfband „{header.get('title', '')}\", Stil "
                     f"{header.get('style', d['headerStyle'])}, Logo "
                     + {"none": "ohne", "right": "rechts"}.get(logo, "links")
                     + (f", Nav: {' · '.join(header.get('nav') or [])}"
                        if header.get("nav") else ""))
    else:
        items.append("Kein Kopfband")
    if z.get("nav"):
        items.append("Linke Nav-Leiste mit Icon je Seite")
    if filt:
        mode = filt.get("mode") or filt.get("side") or "right"
        sl = ", ".join(s_["name"] for s_ in (filt.get("slicers") or [])) or "noch ohne Felder"
        items.append(f"Filter {FILTER_MODE.get(mode, mode)}: {sl}"
                     + (" · ein-/ausblendbar über Lesezeichen"
                        if filt.get("collapsible") or mode == "burger" else ""))
    else:
        items.append("Kein Filter-Panel")
    if z.get("footer"):
        items.append(f"Fußleiste: {z['footer'].get('text', '')}")
    items.append(f"Kacheln {'gerundet (%d px)' % d['cornerRadius'] if d['cornerRadius'] else 'eckig'}"
                 f", {TILE_STYLE.get(d['tileStyle'], d['tileStyle'])}, "
                 f"Hintergrund {d['tileBackground']}")
    items.append(f"Seitenhintergrund {d['pageBackground']} · Akzent {d['accent']} · "
                 f"Schriftfaktor ×{d['fontScale']}")
    s = slide_frame("Entscheidungen zur Gestaltung")
    bullets(s, items)

    # ---------------------------------------------------------- Seiten
    def wireframe(slide, page):
        """Zeichnet die Zonen und Kacheln maßstäblich in die Folie."""
        cw = int(n["canvas"].get("width") or 1280)
        ch = int(n["canvas"].get("height") or 720)
        margin_x, top = 700000, 1450000
        avail_w = SW - 2 * margin_x
        avail_h = SH - top - 700000
        scale = min(avail_w / cw, avail_h / ch)
        ox = int((SW - cw * scale) / 2)
        oy = int(top)

        def box(rx, ry, rw, rh, *, fill, outline=None, text="", sub="",
                size=9, color=ink, bold=False):
            shp = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Emu(int(ox + rx * scale)), Emu(int(oy + ry * scale)),
                Emu(max(9144, int(rw * scale))), Emu(max(9144, int(rh * scale))))
            shp.fill.solid()
            shp.fill.fore_color.rgb = fill
            if outline is None:
                shp.line.fill.background()
            else:
                shp.line.color.rgb = outline
                shp.line.width = Pt(0.75)
            shp.shadow.inherit = False
            tf = shp.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.TOP
            p0 = tf.paragraphs[0]
            p0.alignment = PP_ALIGN.LEFT
            r = p0.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            if sub:
                p1 = tf.add_paragraph()
                p1.alignment = PP_ALIGN.LEFT
                r1 = p1.add_run()
                r1.text = sub
                r1.font.size = Pt(max(7, size - 2))
                r1.font.color.rgb = ink_soft
            return shp

        # Seitenfläche
        box(0, 0, cw, ch, fill=page_bg, outline=line)
        zz = n["zones"]
        if zz.get("nav"):
            r = zz["nav"]
            box(r["x"], r["y"], r["w"], r["h"], fill=ink, text="Nav",
                size=8, color=RGBColor.from_string("FFFFFF"))
        if zz.get("header"):
            r = zz["header"]
            style = r.get("style", d["headerStyle"])
            fill = (RGBColor.from_string("FFFFFF") if style == "light"
                    else accent if style == "accent" else ink)
            fg = ink if style == "light" else RGBColor.from_string("FFFFFF")
            nav = " · ".join(r.get("nav") or [])
            box(r["x"], r["y"], r["w"], r["h"], fill=fill,
                outline=line if style == "light" else None,
                text=r.get("title", "Kopfband"), sub=nav, size=9, color=fg, bold=True)
        if zz.get("filter"):
            r = zz["filter"]
            mode = r.get("mode") or r.get("side") or "right"
            sl = ", ".join(x["name"] for x in (r.get("slicers") or []))
            box(r["x"], r["y"], r["w"], r["h"],
                fill=RGBColor.from_string("EDEFF2"), outline=line,
                text=f"Filter ({mode})", sub=sl, size=8)
        if zz.get("footer"):
            r = zz["footer"]
            box(r["x"], r["y"], r["w"], r["h"], fill=page_bg, outline=line,
                text=zz["footer"].get("text", ""), size=7, color=ink_soft)
        for i, v in enumerate(page["visuals"], start=1):
            r = v.get("rect") or {}
            if not r:
                continue
            sub = (v.get("label") or v.get("kind", ""))
            if v.get("engine") == "ck":
                sub += " · ChartKitchen"
            elif v.get("engine") == "deneb":
                sub += " · Deneb"
            box(r["x"], r["y"], r["w"], r["h"], fill=tile_bg, outline=line,
                text=f"{page['index']}.{i}  {v.get('title', '')}", sub=sub,
                size=9, bold=True)

    for page in n["pages"]:
        s = slide_frame(f"Seite {page['index']} · {page['name']}",
                        page["notes"][:90] if page["notes"] else "")
        wireframe(s, page)

        # Kachel-Tabelle
        s = slide_frame(f"Seite {page['index']} · {page['name']} — Kacheln")
        rows = len(page["visuals"]) + 1
        if rows == 1:
            textbox(s, Emu(457200), Emu(1400000), SW - Emu(914400), Emu(400000),
                    "Noch keine Kacheln auf dieser Seite.", size=14, color=ink_soft)
            continue
        left, top = Emu(457200), Emu(1350000)
        width, height = SW - Emu(914400), Emu(min(4600000, 340000 * rows))
        table = s.shapes.add_table(rows, 4, left, top, width, height).table
        table.columns[0].width = Emu(3100000)
        table.columns[1].width = Emu(2500000)
        table.columns[2].width = Emu(3600000)
        table.columns[3].width = Emu(width - Emu(3100000) - Emu(2500000) - Emu(3600000))
        heads = ["Kachel", "Darstellung", "Felder", "Notizen"]
        for c, head in enumerate(heads):
            cell = table.cell(0, c)
            cell.text = head
            para = cell.text_frame.paragraphs[0]
            para.runs[0].font.size = Pt(11)
            para.runs[0].font.bold = True
            para.runs[0].font.color.rgb = ink
        for r_i, v in enumerate(page["visuals"], start=1):
            render = v.get("label") or v.get("kind", "")
            if v.get("scenario"):
                render += f", {v['scenario']}"
            if v.get("engine") != "native":
                render += f" · {ENGINE_LABEL.get(v.get('engine'), v.get('engine'))}"
            note = " · ".join(x for x in [
                v.get("notes") or "",
                f"→ {v['link']['pageName']}" if v.get("link") else ""] if x)
            title = v.get("title", "")
            if v.get("subtitle"):
                title += f"\n{v['subtitle']}"
            for c, txt in enumerate([f"{page['index']}.{r_i}  {title}", render,
                                     fields_summary(v), note or "–"]):
                cell = table.cell(r_i, c)
                cell.text = str(txt)
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)
                        run.font.color.rgb = ink if c == 0 else ink_soft

    # ---------------------------------------------------------- Navigation
    s = slide_frame("Navigation und Drill-Wege")
    items = []
    if header and header.get("nav"):
        items.append("Nav-Buttons im Kopfband: " + " · ".join(header["nav"])
                     + " (auf jeder Seite, aktive Seite als Akzent-Button ohne Aktion)")
    if z.get("nav"):
        items.append("Linke Nav-Leiste mit einem Button je Seite")
    for l in n["links"]:
        kind = ("Seitenwechsel per Button-Aktion" if l.get("kind") == "navigation"
                else "Drill-through auf das Kategorie-Feld der Quellkachel")
        items.append(f"„{l['fromPage']}\" ({l['fromVisual']}) → „{l['toPage']}\": {kind}")
    if filt and (filt.get("collapsible") or (filt.get("mode") == "burger")):
        items.append("Filter-Panel über zwei Lesezeichen (öffnen / schließen); "
                     "Lesezeichen erfassen nur die Sichtbarkeit, nicht die Auswahl")
    if not items:
        items = ["Keine Navigations- oder Drill-Wege im Mockup festgelegt."]
    bullets(s, items)

    # ---------------------------------------------------------- Kennzahlen
    new = spec.get("newFields") or []
    s = slide_frame("Neue Kennzahlen und Felder")
    if not new:
        textbox(s, Emu(457200), Emu(1400000), SW - Emu(914400), Emu(400000),
                "Keine neuen Felder — alles kommt aus dem bestehenden Modell.",
                size=15, color=ink_soft)
    else:
        rows = len(new) + 1
        table = s.shapes.add_table(rows, 4, Emu(457200), Emu(1350000),
                                   SW - Emu(914400),
                                   Emu(min(4600000, 400000 * rows))).table
        table.columns[0].width = Emu(2600000)
        table.columns[1].width = Emu(1500000)
        table.columns[2].width = Emu(4000000)
        table.columns[3].width = Emu(SW - Emu(914400) - Emu(2600000) - Emu(1500000) - Emu(4000000))
        for c, head in enumerate(["Feld", "Art", "Beschreibung / Logik", "Offene Frage"]):
            cell = table.cell(0, c)
            cell.text = head
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size, run.font.bold, run.font.color.rgb = Pt(11), True, ink
        for i, f in enumerate(new, start=1):
            vals = [f"{f['table']}.{f['name']}",
                    "Kennzahl" if f.get("kind") == "measure" else "Dimension",
                    f.get("description") or "[ausfüllen]",
                    f.get("openQuestion") or "–"]
            for c, txt in enumerate(vals):
                cell = table.cell(i, c)
                cell.text = str(txt)
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(10)
                        run.font.color.rgb = ink if c == 0 else ink_soft

    # ---------------------------------------------------------- Offene Punkte
    op = open_points(spec, n)
    s = slide_frame("Offene Punkte")
    bullets(s, op or ["Keine offenen Punkte erfasst."], size=14, gap=400000)

    # ---------------------------------------------------------- Nächste Schritte
    s = slide_frame("Nächste Schritte")
    bullets(s, [
        "Fehlende Kennzahlen im Semantikmodell anlegen (te add, DAX bestätigen)",
        "Seiten mit dem Skill `mockup-to-powerbi` ins PBIP übertragen",
        "ChartKitchen-Slots über die Referenz-Instanz füllen"
        if any(v.get("engine") == "ck" for p in n["pages"] for v in p["visuals"])
        else "Formatierung gegen das Theme prüfen",
        "Navigation, Drill-through und Filter-Lesezeichen in Desktop testen",
        "Review der gebauten Seiten mit den Teilnehmenden",
    ], size=15)

    target.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(target))
    return True, f"{target}  ({len(prs.slides)} Folien)"


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #
def main() -> int:
    p = argparse.ArgumentParser(
        description="MockupKitchen-Spec → Workshop-Doku (Markdown + PowerPoint)")
    p.add_argument("spec", help="Pfad zu mockup-spec.json")
    p.add_argument("--out", default=None,
                   help="Ausgabeordner (Default: Ordner der Spec)")
    p.add_argument("--docs", default=None,
                   help="vorhandene WORKSHOP-DOKU.md (Default: neben der Spec)")
    p.add_argument("--force-md", action="store_true",
                   help="WORKSHOP-DOKU.md auch dann schreiben, wenn sie schon existiert")
    p.add_argument("--no-pptx", action="store_true", help="nur Markdown erzeugen")
    opt = p.parse_args()

    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

    spec_path = Path(opt.spec).expanduser()
    if not spec_path.is_file():
        print(f"Spec nicht gefunden: {spec_path}", file=sys.stderr)
        return 2
    try:
        spec = json.loads(spec_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        print(f"mockup-spec.json ist kein gültiges JSON: {e}", file=sys.stderr)
        return 2
    if "zones" not in spec or "canvas" not in spec:
        print("Das sieht nicht nach einer MockupKitchen-Spec aus "
              "(zones/canvas fehlen).", file=sys.stderr)
        return 2

    n = normalize(spec)
    out = Path(opt.out).expanduser() if opt.out else spec_path.parent
    out.mkdir(parents=True, exist_ok=True)

    md_path = Path(opt.docs).expanduser() if opt.docs else (out / "WORKSHOP-DOKU.md")
    if md_path.is_file() and not opt.force_md:
        print(f"WORKSHOP-DOKU.md vorhanden, bleibt unverändert: {md_path}")
    else:
        md_path.write_text(build_docs_md(spec, n), encoding="utf-8")
        print(f"WORKSHOP-DOKU.md geschrieben: {md_path}")

    rc = 0
    if opt.no_pptx:
        print("PowerPoint übersprungen (--no-pptx).")
    else:
        ok, msg = build_pptx(spec, n, out / "WORKSHOP-DOKU.pptx")
        print(("PowerPoint: " if ok else "") + msg)
        if not ok:
            rc = 1

    print(f"\n{len(n['pages'])} Seite(n) · "
          f"{sum(len(p['visuals']) for p in n['pages'])} Kacheln · "
          f"{len(spec.get('newFields') or [])} neue Felder · "
          f"{len(open_points(spec, n))} offene Punkte")
    return rc


if __name__ == "__main__":
    raise SystemExit(main())
